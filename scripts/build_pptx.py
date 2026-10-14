#!/usr/bin/env python3
"""Extended manifest -> editable PPTX builder for the tech flowchart slide.

Extends the bggg-creator-image2ppt manifest idea with: any MSO autoshape,
rotation, fill alpha, freeform polygons/polylines and multi-paragraph text.
Coordinates are canvas pixels (default 1920x1080) mapped to a 13.333x7.5in slide.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

DEFAULT_FONT = "Microsoft YaHei"


def px_color(value):
    if not value:
        return None
    text = str(value).strip().lstrip("#")
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def set_typeface(run, name):
    run.font.name = name
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = r_pr.makeelement(qn(tag), {})
            r_pr.append(node)
        node.set("typeface", name)


def set_fill_alpha(shape, alpha_pct):
    """alpha_pct: 0-100 opacity."""
    fill = shape.fill
    srgb = fill.fore_color._xFill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for child in list(srgb):
        srgb.remove(child)
    alpha = srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha_pct * 1000))})
    srgb.append(alpha)


class Scale:
    def __init__(self, canvas_w, canvas_h, slide_w_in, slide_h_in):
        self.cw, self.ch = canvas_w, canvas_h
        self.sw, self.sh = slide_w_in, slide_h_in

    def x(self, v):
        return Inches(float(v) * self.sw / self.cw)

    def y(self, v):
        return Inches(float(v) * self.sh / self.ch)

    def w(self, v):
        return Inches(float(v) * self.sw / self.cw)

    def h(self, v):
        return Inches(float(v) * self.sh / self.ch)

    def lw(self, v):
        # line width: use horizontal ratio, px -> Emu
        return Emu(int(float(v) * 914400 * self.sw / self.cw))


def style_shape(shape, el, sc):
    shape.shadow.inherit = False
    fill = el.get("fill")
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = px_color(fill)
        if el.get("alpha") is not None:
            set_fill_alpha(shape, el["alpha"])
    else:
        shape.fill.background()
    stroke = el.get("stroke")
    if stroke:
        shape.line.color.rgb = px_color(stroke)
        shape.line.width = sc.lw(el.get("stroke_w", 2))
        if el.get("dash"):
            shape.line.dash_style = el["dash"]
    else:
        shape.line.fill.background()
    if el.get("rot"):
        shape.rotation = float(el["rot"])
    if el.get("adj") is not None:
        try:
            shape.adjustments[0] = float(el["adj"])
        except Exception:
            pass


def add_shape_el(slide, el, sc):
    name = str(el.get("shape", "RECTANGLE")).upper()
    mso = getattr(MSO_AUTO_SHAPE_TYPE, name, MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    shape = slide.shapes.add_shape(mso, sc.x(el.get("x", 0)), sc.y(el.get("y", 0)),
                                   sc.w(el.get("w", 10)), sc.h(el.get("h", 10)))
    shape.name = el.get("name", name)
    style_shape(shape, el, sc)
    return shape


def add_line_el(slide, el, sc):
    conn = slide.shapes.add_connector(1, sc.x(el["x1"]), sc.y(el["y1"]), sc.x(el["x2"]), sc.y(el["y2"]))
    conn.name = el.get("name", "Line")
    conn.shadow.inherit = False
    conn.line.color.rgb = px_color(el.get("stroke", "#333333"))
    conn.line.width = sc.lw(el.get("stroke_w", 2))
    return conn


def add_freeform_el(slide, el, sc):
    pts = el["pts"]
    x0, y0 = el.get("x", 0), el.get("y", 0)
    w, h = el.get("w", 100), el.get("h", 100)
    abs_pts = [(sc.x(x0 + px * w), sc.y(y0 + py * h)) for px, py in pts]
    builder = slide.shapes.build_freeform(abs_pts[0][0], abs_pts[0][1], scale=1)
    builder.add_line_segments([(px, py) for px, py in abs_pts[1:]], close=bool(el.get("closed", True)))
    shape = builder.convert_to_shape()
    shape.name = el.get("name", "Freeform")
    style_shape(shape, el, sc)
    return shape


def add_text_el(slide, el, sc):
    box = slide.shapes.add_textbox(sc.x(el.get("x", 0)), sc.y(el.get("y", 0)),
                                   sc.w(el.get("w", 100)), sc.h(el.get("h", 40)))
    box.name = el.get("name", "Text")
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = bool(el.get("wrap", True))
    tf.margin_left = Emu(int(sc.lw(el.get("margin", 2))))
    tf.margin_right = tf.margin_left
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                          "bottom": MSO_ANCHOR.BOTTOM}.get(el.get("valign", "middle"), MSO_ANCHOR.MIDDLE)
    paras = el.get("paras") or [{"text": el.get("text", "")}]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                       "right": PP_ALIGN.RIGHT}.get(para.get("align", el.get("align", "left")), PP_ALIGN.LEFT)
        p.space_after = Pt(float(para.get("space_after", 0)))
        p.space_before = Pt(float(para.get("space_before", 0)))
        if para.get("line_spacing"):
            p.line_spacing = float(para["line_spacing"])
        run = p.add_run()
        run.text = str(para.get("text", ""))
        run.font.size = Pt(float(para.get("size", el.get("size", 12))))
        run.font.bold = bool(para.get("bold", el.get("bold", False)))
        run.font.italic = bool(para.get("italic", el.get("italic", False)))
        run.font.color.rgb = px_color(para.get("color", el.get("color", "#1A1A1A")))
        set_typeface(run, para.get("font", el.get("font", DEFAULT_FONT)))
    return box


def build(manifest_path, output_path):
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    deck = manifest.get("deck", {})
    cw = float(deck.get("canvas_width", 1920))
    ch = float(deck.get("canvas_height", 1080))
    sw = float(deck.get("slide_width_in", 13.333))
    sh = float(deck.get("slide_height_in", sw * ch / cw))
    sc = Scale(cw, ch, sw, sh)
    prs = Presentation()
    prs.slide_width = Inches(sw)
    prs.slide_height = Inches(sh)
    counts = {"shape": 0, "line": 0, "freeform": 0, "text": 0, "background": 0}
    for slide_data in manifest.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for el in slide_data.get("elements", []):
            kind = el.get("kind", "shape")
            if kind == "background":
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = px_color(el.get("fill", "#FFFFFF"))
                counts["background"] += 1
            elif kind == "shape":
                add_shape_el(slide, el, sc)
                counts["shape"] += 1
            elif kind == "line":
                add_line_el(slide, el, sc)
                counts["line"] += 1
            elif kind == "freeform":
                add_freeform_el(slide, el, sc)
                counts["freeform"] += 1
            elif kind == "text":
                add_text_el(slide, el, sc)
                counts["text"] += 1
            else:
                raise ValueError(f"unsupported kind {kind}")
    prs.core_properties.title = deck.get("name", "tech flowchart")
    prs.core_properties.subject = "Generated by bggg-creator-image2ppt project (extended builder)"
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return {"output": str(out.resolve()), "counts": counts,
            "slide_width_in": sw, "slide_height_in": sh,
            "canvas": [cw, ch], "shapes_total": sum(counts.values())}


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()
    print(json.dumps(build(args.manifest, args.output), ensure_ascii=False, indent=2))
